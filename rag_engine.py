import os
import io
from typing import List, Dict, Any, Optional, Tuple
# pyrefly: ignore [missing-import]
from langchain_core.documents import Document
# pyrefly: ignore [missing-import]
from langchain_core.embeddings import Embeddings
# pyrefly: ignore [missing-import]
from langchain_text_splitters import RecursiveCharacterTextSplitter
# pyrefly: ignore [missing-import]
from langchain_community.vectorstores import Chroma
# pyrefly: ignore [missing-import]
from chromadb.utils import embedding_functions


# ==========================================
# 1. ROBUST LOCAL EMBEDDINGS (100% Free & Fast ONNX)
# ==========================================
class LocalONNXEmbeddings(Embeddings):
    """
    Production-ready local embedding wrapper using Chroma's built-in ONNX runtime
    and all-MiniLM-L6-v2. Completely free, runs locally on CPU, and bypasses
    Windows DLL security blocks that affect older libraries.
    """
    def __init__(self):
        self._fn = embedding_functions.DefaultEmbeddingFunction()

    def embed_documents(self, texts: List[str]) -> List[List[float]]:
        # DefaultEmbeddingFunction returns list of float arrays
        embeddings = self._fn(texts)
        return [list(map(float, emb)) for emb in embeddings]

    def embed_query(self, text: str) -> List[float]:
        embedding = self._fn([text])[0]
        return list(map(float, embedding))


# ==========================================
# 2. MULTI-FORMAT DOCUMENT LOADER
# ==========================================
def extract_text_from_file(file_bytes_or_path, filename: str) -> List[Document]:
    """
    Extracts text from PDF, DOCX, TXT, and Markdown files.
    Accepts either a file path (str) or raw bytes / file-like object.
    Returns a list of LangChain Document objects with metadata.
    """
    ext = os.path.splitext(filename)[1].lower()
    documents = []

    # Read binary content
    if isinstance(file_bytes_or_path, (str, os.PathLike)):
        with open(file_bytes_or_path, "rb") as f:
            content_bytes = f.read()
    elif hasattr(file_bytes_or_path, "read"):
        content_bytes = file_bytes_or_path.read()
    elif isinstance(file_bytes_or_path, bytes):
        content_bytes = file_bytes_or_path
    else:
        raise ValueError(f"Unsupported file source type: {type(file_bytes_or_path)}")

    # 1. PDF Parsing
    if ext == ".pdf":
        try:
            # pyrefly: ignore [missing-import]
            import fitz  # PyMuPDF
            doc = fitz.open(stream=content_bytes, filetype="pdf")
            for page_num in range(len(doc)):
                page = doc.load_page(page_num)
                text = page.get_text()
                if text.strip():
                    documents.append(
                        Document(
                            page_content=text,
                            metadata={"source": filename, "page": page_num + 1}
                        )
                    )
            doc.close()
        except Exception:
            # Fallback to pypdf
            # pyrefly: ignore [missing-import]
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(content_bytes))
            for page_num, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                if text.strip():
                    documents.append(
                        Document(
                            page_content=text,
                            metadata={"source": filename, "page": page_num + 1}
                        )
                    )

    # 2. DOCX Parsing
    elif ext in [".docx", ".doc"]:
        # pyrefly: ignore [missing-import]
        import docx
        doc = docx.Document(io.BytesIO(content_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        full_text = "\n\n".join(paragraphs)
        if full_text.strip():
            documents.append(
                Document(
                    page_content=full_text,
                    metadata={"source": filename, "page": 1}
                )
            )

    # 3. Plain Text / Markdown / Code
    else:
        try:
            text = content_bytes.decode("utf-8")
        except UnicodeDecodeError:
            text = content_bytes.decode("latin-1", errors="replace")

        if text.strip():
            documents.append(
                Document(
                    page_content=text,
                    metadata={"source": filename, "page": 1}
                )
            )

    if not documents:
        raise ValueError(f"Could not extract any readable text from '{filename}'. Please ensure the document is not empty or scanned/image-only.")

    return documents


# ==========================================
# 3. TEXT CHUNKING (Splitting)
# ==========================================
def chunk_documents(documents: List[Document], chunk_size: int = 800, chunk_overlap: int = 150) -> List[Document]:
    """
    Splits documents into smaller semantic chunks to ensure proper embedding
    granularity and fit within LLM context windows.
    """
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        length_function=len,
        separators=["\n\n", "\n", ". ", " ", ""]
    )
    chunks = splitter.split_documents(documents)
    # Ensure chunk IDs are recorded in metadata
    for idx, chunk in enumerate(chunks):
        chunk.metadata["chunk_id"] = idx + 1
    return chunks


# ==========================================
# 4. VECTOR STORE CREATION
# ==========================================
def create_vector_store(chunks: List[Document], embeddings: Optional[Embeddings] = None) -> Chroma:
    """
    Indexes document chunks into an in-memory Chroma vector database.
    Uses LocalONNXEmbeddings by default for zero-cost, high-speed semantic retrieval.
    """
    if embeddings is None:
        embeddings = LocalONNXEmbeddings()

    vector_store = Chroma.from_documents(
        documents=chunks,
        embedding=embeddings
    )
    return vector_store


# ==========================================
# 5. LLM FACTORY
# ==========================================
def get_llm(provider: str, api_key: str, model_name: Optional[str] = None, temperature: float = 0.1):
    """
    Returns an LLM client based on selected provider.
    Supports OpenRouter, Google Gemini, OpenAI, and Groq.
    """
    provider = provider.lower().strip()
    if provider == "openrouter":
        # pyrefly: ignore [missing-import]
        from langchain_openrouter import ChatOpenRouter
        selected_model = model_name or "meta-llama/llama-3.2-3b-instruct:free"
        return ChatOpenRouter(
            model=selected_model,
            api_key=api_key,
            temperature=temperature
        )

    elif provider in ["google", "gemini"]:
        # pyrefly: ignore [missing-import]
        from langchain_google_genai import ChatGoogleGenerativeAI
        selected_model = model_name or "gemini-2.5-flash"
        return ChatGoogleGenerativeAI(
            model=selected_model,
            google_api_key=api_key,
            temperature=temperature
        )

    elif provider == "openai":
        # pyrefly: ignore [missing-import]
        from langchain_openai import ChatOpenAI
        selected_model = model_name or "gpt-4o-mini"
        return ChatOpenAI(
            model=selected_model,
            api_key=api_key,
            temperature=temperature
        )

    elif provider == "groq":
        # pyrefly: ignore [missing-import]
        from langchain_openai import ChatOpenAI
        selected_model = model_name or "llama-3.1-8b-instant"
        return ChatOpenAI(
            base_url="https://api.groq.com/openai/v1",
            api_key=api_key,
            model=selected_model,
            temperature=temperature
        )

    else:
        raise ValueError(f"Unsupported provider: {provider}")


# ==========================================
# 6. RAG RETRIEVAL & ANSWER GENERATION
# ==========================================
RAG_SYSTEM_PROMPT = """You are an expert Question-Answering assistant adhering strictly to Real-Life Retrieval-Augmented Generation (RAG) principles.

TASK:
Answer the user's question accurately using ONLY the provided document context excerpts below.

CRITICAL INSTRUCTIONS:
1. Strict Grounding: Use ONLY facts directly mentioned in the Context. Never extrapolate, hallucinate, or bring in outside knowledge.
2. If Not Found: If the context does not contain enough information to answer the question confidently, respond strictly with:
   "I cannot find the answer to this question in the provided document."
3. Transparency: When citing facts, mention the relevant Source or Page number from the context headers if applicable.
4. Formatting: Keep your explanation structured, concise, and easy to read with bullet points when appropriate.

=== RETRIEVED CONTEXT ===
{context}

=== USER QUESTION ===
{question}

=== GROUNDED ANSWER ==="""


def query_rag_system(
    vector_store: Chroma,
    llm,
    question: str,
    top_k: int = 4
) -> Tuple[str, List[Document]]:
    """
    Executes the end-to-end RAG pipeline:
    1. Similarity search in vector database for top_k relevant chunks.
    2. Formats retrieved chunks into grounded prompt context.
    3. Prompts LLM to produce a document-grounded answer.
    4. Returns (answer_text, retrieved_sources).
    """
    # 1. Retrieve relevant chunks
    retriever = vector_store.as_retriever(search_kwargs={"k": top_k})
    retrieved_docs = retriever.invoke(question)

    if not retrieved_docs:
        return "I could not find any relevant information in the uploaded document for your query.", []

    # 2. Format context with source metadata
    context_parts = []
    for i, doc in enumerate(retrieved_docs):
        src = doc.metadata.get("source", "Document")
        page = doc.metadata.get("page", 1)
        chunk_id = doc.metadata.get("chunk_id", i + 1)
        header = f"[Chunk #{chunk_id} | File: {src} | Page: {page}]"
        context_parts.append(f"{header}\n{doc.page_content.strip()}")

    formatted_context = "\n\n---\n\n".join(context_parts)

    # 3. Create prompt and invoke LLM
    # pyrefly: ignore [missing-import]
    from langchain_core.prompts import PromptTemplate
    prompt = PromptTemplate(
        template=RAG_SYSTEM_PROMPT,
        input_variables=["context", "question"]
    )
    chain = prompt | llm

    try:
        response = chain.invoke({
            "context": formatted_context,
            "question": question
        })
        answer_text = response.content if hasattr(response, "content") else str(response)
        return answer_text, retrieved_docs
    except Exception as e:
        err_msg = str(e)
        if "401" in err_msg or "User not found" in err_msg or "API_KEY" in err_msg or "Unauthorized" in err_msg:
            fallback_answer = (
                "⚠️ **[Authentication Notice]**: The configured LLM API key returned a `401 Unauthorized / User Not Found` error.\n"
                "To get live generative responses, replace `OPENROUTER_API_KEY` in `.env` with a valid key from https://openrouter.ai/keys or use Google Gemini.\n\n"
                "📌 **[Grounded Evidence Extracted Directly From Your Document Chunks]**:\n"
            )
            for idx, d in enumerate(retrieved_docs, 1):
                clean_snippet = d.page_content.strip()
                fallback_answer += f"\n• **From Match #{idx} (Page {d.metadata.get('page', 1)})**:\n{clean_snippet}\n"
            return fallback_answer, retrieved_docs
        else:
            raise e

